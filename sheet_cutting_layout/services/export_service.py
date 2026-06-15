from __future__ import annotations

import io
import os
import re
from collections.abc import Callable, Mapping, Sequence
from copy import copy as _copy

from openpyxl import load_workbook

_INVALID_SHEET_TITLE_CHARS = re.compile(r"[\\*?:/\[\]]")
_MAX_SHEET_TITLE_LENGTH = 31


def build_cell_map(layout: Mapping[str, object]) -> dict[str, object]:
	"""Map a Sheet Cutting Layout dict to FRM/PRD/15 worksheet cells."""
	cells: dict[str, object] = {
		"B1": _text(layout.get("company")),
		"A5": f"Sheet Cutting Layout No:- {_text(layout.get('layout_code'))}",
		"G5": _part_name_label(layout),
		"N5": f"Part Number:-{_display_part_numbers(layout)}",
		"B6": _text(layout.get("project")),
		"C6": _text(layout.get("project_name")),
		"J7": _text(layout.get("raw_material_item_name")),
		"K8": _num(layout.get("sheet_thickness_mm")),
		"K9": _num(layout.get("sheet_thickness_mm")),
		"L9": _num(layout.get("sheet_width_mm")),
		"M9": _num(layout.get("sheet_length_mm")),
		"T6": _num(layout.get("weight_per_sheet_kg")),
		"K10": _num(layout.get("weight_of_strip_kg")),
		"K11": _num(layout.get("strip_thickness_mm")),
		"L11": _num(layout.get("strip_width_mm")),
		"M11": _num(layout.get("strip_length_mm")),
		"K12": _num(layout.get("parts_per_strip")),
		"K13": _num(layout.get("no_of_strips")),
		"K14": _num(layout.get("parts_per_sheet")),
		"K15": _num(layout.get("gross_weight_per_part_kg")),
		"K16": _num(layout.get("net_weight_per_part_kg")),
		"K17": _num(layout.get("scrap_weight_per_part_kg")),
	}
	cells.update(_bom_table_cells(layout))
	cells.update(_end_piece_detail_cells(layout))
	return cells


def default_template_path() -> str:
	return os.path.join(
		os.path.dirname(__file__),
		"..",
		"templates",
		"iatf",
		"sheet_cutting_layout.xlsx",
	)


def render_workbook_bytes(layout: Mapping[str, object], template_path: str) -> bytes:
	"""Open the FRM/PRD/15 template, apply the cell map, return .xlsx bytes."""
	workbook = load_workbook(template_path)
	worksheet = workbook.active
	_apply_cell_map(worksheet, layout)
	buffer = io.BytesIO()
	workbook.save(buffer)
	return buffer.getvalue()


def walk_layout_tree(layout: object, fetch_child: Callable[[str], object]) -> list[object]:
	"""Return the layout and every descendant child layout, depth-first."""
	visited: set[str] = set()
	return list(_walk_layout_tree(layout, fetch_child, visited))


def unique_sheet_title(base: object, used: set[str]) -> str:
	"""Return an openpyxl-safe worksheet title unique within ``used``."""
	title = _INVALID_SHEET_TITLE_CHARS.sub("_", str(base or "").strip())
	title = title[:_MAX_SHEET_TITLE_LENGTH] or "Sheet"
	if title not in used:
		return title

	for index in range(2, 10000):
		suffix = f" ({index})"
		trimmed = title[: _MAX_SHEET_TITLE_LENGTH - len(suffix)]
		candidate = f"{trimmed}{suffix}"
		if candidate not in used:
			return candidate
	raise ValueError(f"Could not build a unique sheet title from {base!r}")


def build_multi_sheet_workbook(pages: Sequence[tuple[object, Mapping[str, object]]]):
	"""Build one workbook with one templated worksheet per layout page."""
	pages = list(pages)
	if not pages:
		raise ValueError("Cannot build a workbook without at least one layout page")

	used_titles: set[str] = set()
	workbook = None
	for index, (page_title, layout) in enumerate(pages):
		template = load_workbook(default_template_path())
		source_worksheet = template.active
		title = unique_sheet_title(page_title, used_titles)
		used_titles.add(title)

		if index == 0:
			workbook = template
			worksheet = source_worksheet
			worksheet.title = title
		else:
			worksheet = _clone_template_worksheet(workbook, source_worksheet, title)
		_apply_cell_map(worksheet, layout)

	return workbook


def _walk_layout_tree(layout: object, fetch_child: Callable[[str], object], visited: set[str]):
	name = str(getattr(layout, "name", "") or "").strip()
	if name in visited:
		return
	visited.add(name)
	yield layout

	for row in getattr(layout, "end_pieces", []) or []:
		if str(getattr(row, "disposition", "") or "").strip() != "Reuse":
			continue
		child_name = str(getattr(row, "child_layout", "") or "").strip()
		if not child_name or child_name in visited:
			continue
		child = fetch_child(child_name)
		yield from _walk_layout_tree(child, fetch_child, visited)


def _apply_cell_map(worksheet: object, layout: Mapping[str, object]) -> None:
	for coordinate, value in build_cell_map(layout).items():
		worksheet[coordinate] = _excel_safe_cell(value) if value != "" else None


def _excel_safe_cell(value: object) -> object:
	if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
		return f"'{value}"
	return value


def _clone_template_worksheet(workbook: object, source_worksheet: object, title: str) -> object:
	worksheet = workbook.create_sheet(title=title)
	worksheet.sheet_format = _copy(source_worksheet.sheet_format)
	worksheet.page_margins = _copy(source_worksheet.page_margins)
	worksheet.page_setup = _copy(source_worksheet.page_setup)
	worksheet.print_options = _copy(source_worksheet.print_options)
	worksheet.sheet_properties = _copy(source_worksheet.sheet_properties)
	for row in source_worksheet.iter_rows():
		for cell in row:
			target = worksheet.cell(row=cell.row, column=cell.column, value=cell.value)
			if cell.has_style:
				target.font = _copy(cell.font)
				target.border = _copy(cell.border)
				target.fill = _copy(cell.fill)
				target.number_format = cell.number_format
				target.protection = _copy(cell.protection)
				target.alignment = _copy(cell.alignment)
	for merged_range in list(source_worksheet.merged_cells.ranges):
		worksheet.merge_cells(str(merged_range))
	for column_letter, dimension in source_worksheet.column_dimensions.items():
		worksheet.column_dimensions[column_letter].width = dimension.width
	for row_index, dimension in source_worksheet.row_dimensions.items():
		worksheet.row_dimensions[row_index].height = dimension.height
	for image in getattr(source_worksheet, "_images", []):
		cloned_image = _copy(image)
		cloned_image.anchor = _copy(image.anchor)
		worksheet.add_image(cloned_image)
	return worksheet


def _bom_table_cells(layout: Mapping[str, object]) -> dict[str, object]:
	cells: dict[str, object] = {
		coordinate: ""
		for row in (9, 10, 11)
		for coordinate in (f"O{row}", f"Q{row}", f"R{row}", f"S{row}", f"T{row}", f"U{row}")
	}
	gross = layout.get("gross_weight_per_part_kg")
	nos = layout.get("parts_per_sheet")
	cells["O8"] = _primary_part_number(layout)
	cells["Q8"] = _num(gross)
	cells["R8"] = _num(layout.get("net_weight_per_part_kg"))
	cells["S8"] = _num(layout.get("scrap_weight_per_part_kg"))
	cells["T8"] = _num(nos)
	cells["U8"] = _num(_product(gross, nos))

	end_pieces: Sequence[Mapping[str, object]] = layout.get("end_pieces") or []
	for offset, end_piece in enumerate(end_pieces):
		row = 9 + offset
		if row > 10:
			break
		ep_gross = end_piece.get("gross_weight_per_part_kg")
		ep_nos = end_piece.get("bom_quantity")
		cells[f"O{row}"] = _end_piece_label(end_piece)
		cells[f"Q{row}"] = _num(ep_gross)
		cells[f"R{row}"] = _num(end_piece.get("net_weight_per_part_kg"))
		cells[f"S{row}"] = _num(end_piece.get("scrap_weight_per_part_kg"))
		cells[f"T{row}"] = _num(ep_nos)
		cells[f"U{row}"] = _num(_product(ep_gross, ep_nos))
	return cells


def _end_piece_detail_cells(layout: Mapping[str, object]) -> dict[str, object]:
	blocks = (
		{
			"size": ("K18", "L18", "M18"),
			"detail_size": ("K22", "L22", "M22"),
			"used_for": "J23",
			"strip": ("K24", "L24", "M24"),
			"parts": "K25",
			"gross": "K26",
			"net": "K27",
			"scrap": "K28",
		},
		{
			"size": ("R12", "S12", "T12"),
			"detail_size": (),
			"used_for": "Q13",
			"strip": ("R14", "S14", "T14"),
			"parts": "R15",
			"gross": "R16",
			"net": "R17",
			"scrap": "R18",
		},
		{
			"size": ("R22", "S22", "T22"),
			"detail_size": (),
			"used_for": "Q23",
			"strip": ("R24", "S24", "T24"),
			"parts": "R25",
			"gross": "R26",
			"net": "R27",
			"scrap": "R28",
		},
	)
	cells = {
		coordinate: ""
		for block in blocks
		for coordinate in (
			*block["size"],
			*block["detail_size"],
			block["used_for"],
			*block["strip"],
			block["parts"],
			block["gross"],
			block["net"],
			block["scrap"],
		)
	}

	sheet_thickness = layout.get("sheet_thickness_mm")
	end_pieces: Sequence[Mapping[str, object]] = layout.get("end_pieces") or []
	for end_piece, block in zip(end_pieces, blocks, strict=False):
		size_thickness, size_width, size_length = block["size"]
		detail_size_cells = block["detail_size"]
		strip_thickness, strip_width, strip_length = block["strip"]
		cells[size_thickness] = _num(sheet_thickness)
		cells[size_width] = _num(end_piece.get("width_mm"))
		cells[size_length] = _num(end_piece.get("length_mm"))
		if detail_size_cells:
			detail_thickness, detail_width, detail_length = detail_size_cells
			cells[detail_thickness] = _num(sheet_thickness)
			cells[detail_width] = _num(end_piece.get("width_mm"))
			cells[detail_length] = _num(end_piece.get("length_mm"))
		cells[block["used_for"]] = _text(end_piece.get("used_for_finished_part"))
		cells[strip_thickness] = _num(sheet_thickness)
		cells[strip_width] = _num(end_piece.get("width_mm"))
		cells[strip_length] = _num(end_piece.get("length_mm"))
		cells[block["parts"]] = _num(end_piece.get("bom_quantity"))
		cells[block["gross"]] = _num(end_piece.get("gross_weight_per_part_kg"))
		cells[block["net"]] = _num(end_piece.get("net_weight_per_part_kg"))
		cells[block["scrap"]] = _num(end_piece.get("scrap_weight_per_part_kg"))

	return cells


def _part_name_label(layout: Mapping[str, object]) -> str:
	base = _text(layout.get("part_name")).strip()
	if layout.get("is_lh_rh"):
		part_names = _clean_strings(layout.get("part_names") or [])
		if part_names:
			return f"Part Name:-{' & '.join(part_names)}"
		if base:
			return f"Part Name:-{base} LH & RH"
	if not base:
		return ""
	return f"Part Name:-{base}"


def _display_part_numbers(layout: Mapping[str, object]) -> str:
	return "/".join(_clean_strings(layout.get("part_numbers") or []))


def _primary_part_number(layout: Mapping[str, object]) -> str:
	part_numbers = _clean_strings(layout.get("part_numbers") or [])
	return part_numbers[0] if part_numbers else ""


def _clean_strings(values: Sequence[object]) -> list[str]:
	return [str(value).strip() for value in values if str(value or "").strip()]


def _end_piece_label(end_piece: Mapping[str, object]) -> str:
	item_code = _text(end_piece.get("end_piece_item_code")).strip()
	if item_code:
		return item_code
	width = end_piece.get("width_mm")
	length = end_piece.get("length_mm")
	if width and length:
		return f"{width} x {length}"
	return ""


def _num(value: object) -> object:
	if value is None or value == "":
		return ""
	return value


def _text(value: object) -> str:
	if value is None:
		return ""
	return str(value)


def _product(left: object, right: object) -> object:
	if left in (None, "") or right in (None, ""):
		return ""
	try:
		return float(left) * float(right)
	except (TypeError, ValueError):
		return ""
