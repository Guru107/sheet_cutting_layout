from __future__ import annotations

import io
from types import SimpleNamespace

import frappe
from openpyxl import load_workbook

from sheet_cutting_layout.tests.base import SheetCuttingLayoutTestCase
from sheet_cutting_layout.tests.factories import ensure_hsn_code, ensure_item_group


def _ensure_item(item_code: str, item_name: str | None = None) -> str:
	if not frappe.db.exists("Item", item_code):
		doc: dict[str, object] = {
			"doctype": "Item",
			"item_code": item_code,
			"item_name": item_name or item_code,
			"item_group": ensure_item_group(),
			"stock_uom": "Kg",
			"is_stock_item": 1,
		}
		if frappe.get_meta("Item", cached=True).has_field("gst_hsn_code") and frappe.db.exists(
			"DocType", "GST HSN Code"
		):
			doc["gst_hsn_code"] = ensure_hsn_code("720810")
		frappe.get_doc(doc).insert(ignore_permissions=True)
	return item_code


def _ensure_project(project_name: str) -> str:
	existing = frappe.db.exists("Project", {"project_name": project_name})
	if existing:
		return str(existing)
	return (
		frappe.get_doc({"doctype": "Project", "project_name": project_name})
		.insert(ignore_permissions=True)
		.name
	)


class TestSheetCuttingLayoutExport(SheetCuttingLayoutTestCase):
	def _build_layout(
		self,
		*,
		is_lh_rh: bool = False,
		with_end_piece: bool = False,
		with_scrap_end_piece: bool = False,
	) -> str:
		if with_end_piece and with_scrap_end_piece:
			raise ValueError("Choose either with_end_piece or with_scrap_end_piece")

		suffix = frappe.generate_hash(length=8).upper()
		project = _ensure_project(f"SCL Export {suffix}")
		raw_material = _ensure_item(f"SCLRM{suffix}", item_name=f"Raw Material {suffix}")
		scrap = _ensure_item(f"SCLSCRAP{suffix}")
		finished = _ensure_item(f"SCLPART{suffix}SHR", item_name="Brkt bumper top")
		twin = _ensure_item(f"SCLTWIN{suffix}SHR", item_name="Brkt bumper top RH")
		layout_code = f"SCL-EXPORT-LAYOUT-{suffix}"
		layout: dict[str, object] = {
			"doctype": "Sheet Cutting Layout",
			"layout_code": layout_code,
			"project": project,
			"revision_no": 1,
			"workflow_status": "Draft",
			"raw_material_item": raw_material,
			"process_scrap_item": scrap,
			"sheet_thickness_mm": 2.0,
			"sheet_width_mm": 1000.0,
			"sheet_length_mm": 2000.0,
			"weight_per_sheet_kg": 31.44,
			"strip_thickness_mm": 2.0,
			"strip_width_mm": 1000.0,
			"strip_length_mm": 1000.0,
			"weight_of_strip_kg": 15.72,
			"parts_per_strip": 1,
			"no_of_strips": 2,
			"parts_per_sheet": 2,
			"finished_part_code": finished,
			"net_weight_per_part_kg": 15.52,
			"gross_weight_per_part_kg": 15.72,
			"scrap_weight_per_part_kg": 0.2,
		}
		if with_end_piece:
			layout.update(
				{
					"net_weight_per_part_kg": 15.0485,
					"strip_length_mm": 970.005,
				}
			)
			layout["end_pieces"] = [
				{
					"doctype": "Layout End Piece",
					"width_mm": 200.0,
					"length_mm": 300.0,
					"weight_kg": 0.943,
					"disposition": "Reuse",
					"used_for_finished_part": finished,
					"scrap_item": scrap,
					"bom_quantity": 2,
					"net_weight_per_part_kg": 0.4,
				}
			]
		if with_scrap_end_piece:
			layout.update(
				{
					"net_weight_per_part_kg": 15.0485,
					"strip_length_mm": 970.005,
				}
			)
			layout["end_pieces"] = [
				{
					"doctype": "Layout End Piece",
					"width_mm": 200.0,
					"length_mm": 300.0,
					"disposition": "Scrap",
					"scrap_item": scrap,
				}
			]
		if is_lh_rh:
			layout.update({"is_lh_rh": 1, "orientation": "LH", "twin_finished_part": twin})
		doc = frappe.get_doc(layout)
		doc.insert(ignore_permissions=True)
		return doc.name

	def test_download_produces_workbook_with_expected_cells(self) -> None:
		from sheet_cutting_layout.sheet_cutting_layout.doctype.sheet_cutting_layout.sheet_cutting_layout import (
			download_sheet_cutting_layout,
		)

		layout_name = self._build_layout()
		frappe.response.clear()

		download_sheet_cutting_layout(layout_name)

		self.assertEqual(frappe.response["type"], "binary")
		self.assertTrue(str(frappe.response["filename"]).endswith(".xlsx"))
		content = frappe.response["filecontent"]
		self.assertIsInstance(content, bytes)
		self.assertEqual(content[:2], b"PK")

		worksheet = load_workbook(io.BytesIO(content)).active
		self.assertTrue(str(worksheet["A5"].value).startswith("Sheet Cutting Layout No:- SCL-EXPORT-LAYOUT-"))
		self.assertIn("Raw Material", str(worksheet["J7"].value))
		self.assertEqual(worksheet["K9"].value, 2.0)
		self.assertEqual(worksheet["L9"].value, 1000.0)
		self.assertEqual(worksheet["M9"].value, 2000.0)
		self.assertEqual(worksheet["Q6"].value, 2.0)
		self.assertEqual(worksheet["R6"].value, 1000.0)
		self.assertEqual(worksheet["S6"].value, 2000.0)
		self.assertEqual(worksheet["T6"].value, 31.44)
		self.assertEqual(worksheet["K14"].value, 2)
		self.assertEqual(worksheet["K15"].value, 15.72)
		self.assertEqual(worksheet["K16"].value, 15.52)
		self.assertIn("SCL Export", str(worksheet["B6"].value))
		self.assertIsNone(worksheet["C6"].value)
		self.assertIn("SCLPART", str(worksheet["N5"].value))
		self.assertTrue(str(worksheet["N5"].value).startswith("Part Number:-"))
		self.assertEqual(worksheet["G5"].value, "Part Name:-Brkt bumper top")
		self.assertNotEqual(worksheet["G5"].value, worksheet["N5"].value)
		self.assertEqual(worksheet["U8"].value, 31.44)
		self.assertIn(worksheet["O9"].value, (None, ""))
		self.assertIn(worksheet["Q9"].value, (None, ""))
		self.assertIn(worksheet["U11"].value, (None, ""))

	def test_download_uses_sheet_cutting_layout_settings_header(self) -> None:
		from sheet_cutting_layout.sheet_cutting_layout.doctype.sheet_cutting_layout.sheet_cutting_layout import (
			download_sheet_cutting_layout,
		)

		settings = frappe.get_single("Sheet Cutting Layout Settings")
		settings.document_number = "SCL/DOC/09"
		settings.revision_number = "04"
		settings.revision_date = "2026-06-23"
		settings.page_text = "01 OF 02"
		settings.save(ignore_permissions=True)

		layout_name = self._build_layout()
		frappe.response.clear()

		download_sheet_cutting_layout(layout_name)

		worksheet = load_workbook(io.BytesIO(frappe.response["filecontent"])).active
		self.assertEqual(worksheet["Q1"].value, "DOC. NO.: SCL/DOC/09")
		self.assertEqual(worksheet["Q2"].value, "REV. NO.: 04")
		self.assertEqual(worksheet["Q3"].value, "REV DATE.: 23.06.2026")
		self.assertEqual(worksheet["Q4"].value, "PAGE: 01 OF 02")

	def test_download_populates_end_piece_detail_from_system_snapshot(self) -> None:
		from sheet_cutting_layout.sheet_cutting_layout.doctype.sheet_cutting_layout.sheet_cutting_layout import (
			download_sheet_cutting_layout,
		)

		layout_name = self._build_layout(with_end_piece=True)
		frappe.response.clear()

		download_sheet_cutting_layout(layout_name)

		worksheet = load_workbook(io.BytesIO(frappe.response["filecontent"])).active
		self.assertAlmostEqual(float(worksheet["Q9"].value), 0.4716, places=4)
		self.assertEqual(worksheet["R9"].value, 0.4)
		self.assertAlmostEqual(float(worksheet["S9"].value), 0.0716, places=4)
		self.assertEqual(worksheet["T9"].value, 2)
		self.assertAlmostEqual(float(worksheet["U9"].value), 0.943, places=3)
		self.assertEqual(worksheet["K18"].value, 2.0)
		self.assertEqual(worksheet["L18"].value, 200.0)
		self.assertEqual(worksheet["M18"].value, 300.0)
		self.assertTrue(str(worksheet["J24"].value).startswith("SCLPART"))
		self.assertEqual(worksheet["K23"].value, 2.0)
		self.assertEqual(worksheet["L23"].value, 200.0)
		self.assertEqual(worksheet["M23"].value, 300.0)
		self.assertEqual(worksheet["K25"].value, 2.0)
		self.assertEqual(worksheet["L25"].value, 200.0)
		self.assertEqual(worksheet["M25"].value, 300.0)
		self.assertEqual(worksheet["K26"].value, 2)
		self.assertAlmostEqual(float(worksheet["K27"].value), 0.4716, places=4)
		self.assertEqual(worksheet["K28"].value, 0.4)
		self.assertAlmostEqual(float(worksheet["K29"].value), 0.0716, places=4)
		self.assertIn(worksheet["O10"].value, (None, ""))
		self.assertIn(worksheet["U10"].value, (None, ""))

	def test_export_end_piece_dict_preserves_reuse_strip_dimensions_for_all_blocks(self) -> None:
		from sheet_cutting_layout.services.export_service import build_multi_sheet_workbook
		from sheet_cutting_layout.sheet_cutting_layout.doctype.sheet_cutting_layout.sheet_cutting_layout import (
			_export_end_piece_dict,
		)

		rows = [
			SimpleNamespace(
				end_piece_item_code=f"FG01SHR-EP-{index}",
				disposition="Reuse",
				used_for_finished_part=f"FG00{index}SHR",
				width_mm=1250.0,
				length_mm=179.0 + index,
				weight_kg=2.81388,
				strip_width_mm=1250.0,
				strip_length_mm=170.0 + index,
				strip_weight_kg=2.6724,
				gross_weight_per_part_kg=0.381771,
				net_weight_per_part_kg=0.171,
				scrap_weight_per_part_kg=0.210771,
				bom_quantity=7,
			)
			for index in range(3)
		]
		exported = [_export_end_piece_dict(row) for row in rows]

		worksheet = build_multi_sheet_workbook(
			[("export", {"sheet_thickness_mm": 1.6, "end_pieces": exported})]
		).active

		self.assertEqual(worksheet["K25"].value, 1.6)
		self.assertEqual(worksheet["L25"].value, 1250.0)
		self.assertEqual(worksheet["M25"].value, 170.0)
		self.assertEqual(worksheet["R15"].value, 1.6)
		self.assertEqual(worksheet["S15"].value, 1250.0)
		self.assertEqual(worksheet["T15"].value, 171.0)
		self.assertEqual(worksheet["R25"].value, 1.6)
		self.assertEqual(worksheet["S25"].value, 1250.0)
		self.assertEqual(worksheet["T25"].value, 172.0)

	def test_download_populates_scrap_end_piece_bom_weight(self) -> None:
		from sheet_cutting_layout.sheet_cutting_layout.doctype.sheet_cutting_layout.sheet_cutting_layout import (
			download_sheet_cutting_layout,
		)

		layout_name = self._build_layout(with_scrap_end_piece=True)
		frappe.response.clear()

		download_sheet_cutting_layout(layout_name)

		worksheet = load_workbook(io.BytesIO(frappe.response["filecontent"])).active
		self.assertEqual(worksheet["O9"].value, "ENDPIECE")
		self.assertAlmostEqual(float(worksheet["Q9"].value), 0.943, places=3)
		self.assertEqual(worksheet["R9"].value, 0)
		self.assertAlmostEqual(float(worksheet["S9"].value), 0.943, places=3)
		self.assertEqual(worksheet["T9"].value, 1)
		self.assertAlmostEqual(float(worksheet["U9"].value), 0.943, places=3)
		self.assertAlmostEqual(float(worksheet["U8"].value) + float(worksheet["U9"].value), 31.44, places=2)

	def test_download_uses_persisted_lh_rh_pairing(self) -> None:
		from sheet_cutting_layout.sheet_cutting_layout.doctype.sheet_cutting_layout.sheet_cutting_layout import (
			download_sheet_cutting_layout,
		)

		layout_name = self._build_layout(is_lh_rh=True)
		frappe.response.clear()

		download_sheet_cutting_layout(layout_name)

		worksheet = load_workbook(io.BytesIO(frappe.response["filecontent"])).active
		self.assertEqual(worksheet["G5"].value, "Part Name:-Brkt bumper top & Brkt bumper top RH")
		self.assertTrue(str(worksheet["N5"].value).startswith("Part Number:-SCLPART"))
		self.assertIn("/SCLTWIN", str(worksheet["N5"].value))
		self.assertNotIn("_SCLTWIN", str(worksheet["N5"].value))
		self.assertTrue(str(worksheet["N5"].value).endswith("SHR"))
