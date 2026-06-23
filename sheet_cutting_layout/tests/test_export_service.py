from __future__ import annotations

import base64
import io
import tempfile

from openpyxl import load_workbook

from sheet_cutting_layout.tests.base import SheetCuttingLayoutTestCase


def _render_workbook_content(layout: dict[str, object]) -> bytes:
	from sheet_cutting_layout.services.export_service import build_multi_sheet_workbook

	stream = io.BytesIO()
	build_multi_sheet_workbook([("Sheet", layout)]).save(stream)
	return stream.getvalue()


def _base_layout() -> dict[str, object]:
	return {
		"company": "Acme Press Parts",
		"layout_code": "SCL-EXPORT-LAYOUT-001",
		"part_name": "Brkt bumper top",
		"part_numbers": ["0102AAG06400SHR"],
		"is_lh_rh": False,
		"project": "PRJ-0001",
		"project_name": "Bumper Program",
		"raw_material_item_name": "HSLA-340",
		"sheet_thickness_mm": 2.0,
		"sheet_width_mm": 1000.0,
		"sheet_length_mm": 2000.0,
		"weight_per_sheet_kg": 31.44,
		"weight_of_strip_kg": 15.72,
		"strip_thickness_mm": 2.0,
		"strip_width_mm": 1000.0,
		"strip_length_mm": 1000.0,
		"parts_per_strip": 1,
		"no_of_strips": 2,
		"parts_per_sheet": 2,
		"gross_weight_per_part_kg": 15.72,
		"net_weight_per_part_kg": 15.52,
		"scrap_weight_per_part_kg": 0.2,
		"raw_material_weight_kg": 31.44,
		"end_pieces": [],
	}


class TestApprovedIatfTemplate(SheetCuttingLayoutTestCase):
	def test_committed_template_matches_approved_visual_anchors(self) -> None:
		from openpyxl import load_workbook

		from sheet_cutting_layout.services.export_service import default_template_path

		workbook = load_workbook(default_template_path())
		self.assertEqual(len(workbook.worksheets), 1)
		worksheet = workbook.active

		self.assertEqual(worksheet.page_setup.orientation, "landscape")
		self.assertEqual(str(worksheet.page_setup.paperSize), "9")
		self.assertIn("A1:A4", {str(item) for item in worksheet.merged_cells.ranges})
		self.assertIn("E1:P4", {str(item) for item in worksheet.merged_cells.ranges})
		self.assertIn("G5:M5", {str(item) for item in worksheet.merged_cells.ranges})
		self.assertIn("N5:Q5", {str(item) for item in worksheet.merged_cells.ranges})
		self.assertIn("R39:U41", {str(item) for item in worksheet.merged_cells.ranges})
		self.assertEqual(worksheet["Q1"].value, "DOC. NO.: ")
		self.assertEqual(len(getattr(worksheet, "_images", [])), 0)
		self.assertEqual(worksheet["O7"].value, "BOM")
		self.assertEqual(worksheet["Q7"].value, "Gross Wt")
		self.assertEqual(worksheet["R7"].value, "F.g Wt")
		self.assertEqual(worksheet["S7"].value, "Scrap Wt")
		self.assertEqual(worksheet["R39"].value, "Released By    \nManagement Rep")


class TestBuildCellMapHeaderAndStrip(SheetCuttingLayoutTestCase):
	def test_header_strip_and_per_part_cells(self) -> None:
		from sheet_cutting_layout.services.export_service import build_cell_map

		cells = build_cell_map(_base_layout())

		self.assertEqual(cells["B1"], "Acme Press Parts")
		self.assertEqual(cells["A5"], "Sheet Cutting Layout No:- SCL-EXPORT-LAYOUT-001")
		self.assertEqual(cells["G5"], "Part Name:-Brkt bumper top")
		self.assertEqual(cells["N5"], "Part Number:-0102AAG06400SHR")
		self.assertEqual(cells["B6"], "Bumper Program")
		self.assertNotIn("C6", cells)
		self.assertEqual(cells["J7"], "HSLA-340")
		self.assertEqual(cells["K8"], 2.0)
		self.assertEqual(cells["K9"], 2.0)
		self.assertEqual(cells["L9"], 1000.0)
		self.assertEqual(cells["M9"], 2000.0)
		self.assertEqual(cells["Q6"], 2.0)
		self.assertEqual(cells["R6"], 1000.0)
		self.assertEqual(cells["S6"], 2000.0)
		self.assertEqual(cells["T6"], 31.44)
		self.assertEqual(cells["K10"], 15.72)
		self.assertEqual(cells["K11"], 2.0)
		self.assertEqual(cells["L11"], 1000.0)
		self.assertEqual(cells["M11"], 1000.0)
		self.assertEqual(cells["K12"], 1)
		self.assertEqual(cells["K13"], 2)
		self.assertEqual(cells["K14"], 2)
		self.assertEqual(cells["K15"], 15.72)
		self.assertEqual(cells["K16"], 15.52)
		self.assertEqual(cells["K17"], 0.2)
		self.assertEqual(cells["O8"], "0102AAG06400SHR")
		self.assertNotIn("P8", cells)
		self.assertEqual(cells["Q8"], 15.72)
		self.assertEqual(cells["R8"], 15.52)
		self.assertEqual(cells["S8"], 0.2)
		self.assertEqual(cells["T8"], 2)
		self.assertAlmostEqual(float(cells["U8"]), 31.44, places=6)


class TestBuildCellMapLhRhLabels(SheetCuttingLayoutTestCase):
	def test_lh_rh_part_name_and_joined_numbers(self) -> None:
		from sheet_cutting_layout.services.export_service import build_cell_map

		layout = _base_layout()
		layout.update(
			{
				"part_name": "Brkt bumper top",
				"part_names": ["Brkt bumper top", "Brkt bumper top RH"],
				"is_lh_rh": True,
				"orientation": "LH",
				"part_numbers": ["0102AAG06400SHR", "0102AAG06410SHR"],
			}
		)

		cells = build_cell_map(layout)

		self.assertEqual(cells["G5"], "Part Name:-Brkt bumper top & Brkt bumper top RH")
		self.assertEqual(cells["N5"], "Part Number:-0102AAG06400SHR/0102AAG06410SHR")
		self.assertEqual(cells["O8"], "0102AAG06400SHR")

	def test_single_part_number_is_not_joined(self) -> None:
		from sheet_cutting_layout.services.export_service import build_cell_map

		layout = _base_layout()
		layout["part_numbers"] = ["0102AAG06400SHR"]

		cells = build_cell_map(layout)

		self.assertEqual(cells["N5"], "Part Number:-0102AAG06400SHR")
		self.assertEqual(cells["G5"], "Part Name:-Brkt bumper top")

	def test_blank_part_numbers_are_dropped_from_join(self) -> None:
		from sheet_cutting_layout.services.export_service import build_cell_map

		layout = _base_layout()
		layout["part_numbers"] = ["0102AAG06400SHR", "", None, "  "]

		cells = build_cell_map(layout)

		self.assertEqual(cells["N5"], "Part Number:-0102AAG06400SHR")


class TestBuildCellMapEndPiecesAndGuards(SheetCuttingLayoutTestCase):
	def test_end_piece_rows_populate_bom_table(self) -> None:
		from sheet_cutting_layout.services.export_service import build_cell_map

		layout = _base_layout()
		layout["end_pieces"] = [
			{
				"disposition": "Reuse",
				"end_piece_item_code": "EP-0102AAG06400-1",
				"used_for_finished_part": "0102AAG06400SHR",
				"width_mm": 200.0,
				"length_mm": 300.0,
				"weight_kg": 0.94,
				"gross_weight_per_part_kg": 0.47,
				"net_weight_per_part_kg": 0.4,
				"scrap_weight_per_part_kg": 0.07,
				"bom_quantity": 2,
			}
		]

		cells = build_cell_map(layout)

		self.assertEqual(cells["O9"], "0102AAG06400SHR")
		self.assertNotIn("P9", cells)
		self.assertEqual(cells["Q9"], 0.47)
		self.assertEqual(cells["R9"], 0.4)
		self.assertEqual(cells["S9"], 0.07)
		self.assertEqual(cells["T9"], 2)
		self.assertAlmostEqual(float(cells["U9"]), 0.94, places=6)
		self.assertEqual(cells["O10"], "")
		self.assertEqual(cells["U10"], "")
		self.assertEqual(cells["O11"], "")
		self.assertEqual(cells["U11"], "")

	def test_scrap_end_piece_rows_populate_bom_table_with_weight(self) -> None:
		from sheet_cutting_layout.services.export_service import build_cell_map

		layout = _base_layout()
		layout["end_pieces"] = [
			{
				"disposition": "Scrap",
				"width_mm": 200.0,
				"length_mm": 300.0,
				"weight_kg": 0.94,
			},
			{
				"disposition": "Scrap",
				"width_mm": 100.0,
				"length_mm": 300.0,
				"weight_kg": 0.47,
			},
			{
				"disposition": "Scrap",
				"width_mm": 50.0,
				"length_mm": 300.0,
				"weight_kg": 0.235,
			},
		]

		cells = build_cell_map(layout)

		self.assertEqual(cells["O9"], "ENDPIECE")
		self.assertEqual(cells["Q9"], 0.94)
		self.assertEqual(cells["R9"], 0)
		self.assertEqual(cells["S9"], 0.94)
		self.assertEqual(cells["T9"], 1)
		self.assertEqual(cells["U9"], 0.94)
		self.assertEqual(cells["O10"], "ENDPIECE")
		self.assertEqual(cells["U10"], 0.47)
		self.assertEqual(cells["O11"], "ENDPIECE")
		self.assertEqual(cells["U11"], 0.235)

	def test_non_reuse_end_piece_uses_generic_label(self) -> None:
		from sheet_cutting_layout.services.export_service import build_cell_map

		layout = _base_layout()
		layout["end_pieces"] = [
			{
				"disposition": "Scrap",
				"width_mm": 200.0,
				"length_mm": 300.0,
				"strip_width_mm": 100.0,
				"strip_length_mm": 150.0,
				"weight_kg": 0.94,
			}
		]

		cells = build_cell_map(layout)

		self.assertEqual(cells["O9"], "ENDPIECE")
		self.assertEqual(cells["L25"], 200.0)
		self.assertEqual(cells["M25"], 300.0)

	def test_first_three_end_pieces_render_in_bom_and_detail_blocks(self) -> None:
		from sheet_cutting_layout.services.export_service import build_cell_map

		layout = _base_layout()
		layout["end_pieces"] = [
			{
				"disposition": "Reuse",
				"end_piece_item_code": f"EP-{index}",
				"width_mm": 200 + index,
				"length_mm": 300 + index,
				"strip_width_mm": 100 + index,
				"strip_length_mm": 150 + index,
				"gross_weight_per_part_kg": 0.4 + index,
				"net_weight_per_part_kg": 0.3 + index,
				"scrap_weight_per_part_kg": 0.1 + index,
				"bom_quantity": index + 2,
				"used_for_finished_part": f"0102AAG0640{index}SHR",
			}
			for index in range(5)
		]

		cells = build_cell_map(layout)

		self.assertEqual(cells["O9"], "0102AAG06400SHR")
		self.assertEqual(cells["O10"], "0102AAG06401SHR")
		self.assertEqual(cells["O11"], "0102AAG06402SHR")
		self.assertNotIn("P10", cells)
		self.assertEqual(cells["K18"], 2.0)
		self.assertEqual(cells["L18"], 200)
		self.assertEqual(cells["M18"], 300)
		self.assertEqual(cells["J24"], "0102AAG06400SHR")
		self.assertEqual(cells["K23"], 2.0)
		self.assertEqual(cells["L23"], 200)
		self.assertEqual(cells["M23"], 300)
		self.assertEqual(cells["K25"], 2.0)
		self.assertEqual(cells["L25"], 100)
		self.assertEqual(cells["M25"], 150)
		self.assertEqual(cells["K26"], 2)
		self.assertEqual(cells["K27"], 0.4)
		self.assertEqual(cells["K28"], 0.3)
		self.assertEqual(cells["K29"], 0.1)
		self.assertEqual(cells["K19"], 2.0)
		self.assertEqual(cells["L19"], 201)
		self.assertEqual(cells["M19"], 301)
		self.assertEqual(cells["Q14"], "0102AAG06401SHR")
		self.assertEqual(cells["R15"], 2.0)
		self.assertEqual(cells["S15"], 101)
		self.assertEqual(cells["T15"], 151)
		self.assertEqual(cells["R16"], 3)
		self.assertEqual(cells["R17"], 1.4)
		self.assertEqual(cells["R18"], 1.3)
		self.assertEqual(cells["R19"], 1.1)
		self.assertEqual(cells["K20"], 2.0)
		self.assertEqual(cells["L20"], 202)
		self.assertEqual(cells["M20"], 302)
		self.assertEqual(cells["Q24"], "0102AAG06402SHR")
		self.assertEqual(cells["R25"], 2.0)
		self.assertEqual(cells["S25"], 102)
		self.assertEqual(cells["T25"], 152)
		self.assertEqual(cells["R26"], 4)
		self.assertEqual(cells["R27"], 2.4)
		self.assertEqual(cells["R28"], 2.3)
		self.assertEqual(cells["R29"], 2.1)
		self.assertNotIn("EP-3", cells.values())
		self.assertNotIn("EP-4", cells.values())

	def test_absent_end_piece_detail_blocks_clear_template_cells(self) -> None:
		from sheet_cutting_layout.services.export_service import build_cell_map

		cells = build_cell_map(_base_layout())

		for coordinate in (
			"K18",
			"L18",
			"M18",
			"K23",
			"L23",
			"M23",
			"J24",
			"K25",
			"K26",
			"K27",
			"K28",
			"K29",
			"K19",
			"L19",
			"M19",
			"R13",
			"S13",
			"T13",
			"Q14",
			"R15",
			"R16",
			"R17",
			"R18",
			"R19",
			"K20",
			"L20",
			"M20",
			"R23",
			"S23",
			"T23",
			"Q24",
			"R25",
			"R26",
			"R27",
			"R28",
			"R29",
		):
			self.assertEqual(cells[coordinate], "")

	def test_unused_bom_rows_clear_template_placeholders(self) -> None:
		from sheet_cutting_layout.services.export_service import build_cell_map

		cells = build_cell_map(_base_layout())

		for row in (9, 10, 11):
			for column in ("O", "Q", "R", "S", "T", "U"):
				self.assertEqual(cells[f"{column}{row}"], "")

	def test_missing_numeric_inputs_render_as_empty_string(self) -> None:
		from sheet_cutting_layout.services.export_service import build_cell_map

		layout = _base_layout()
		layout.update(
			{
				"sheet_thickness_mm": None,
				"weight_of_strip_kg": None,
				"gross_weight_per_part_kg": None,
				"net_weight_per_part_kg": "",
				"parts_per_sheet": None,
				"weight_per_sheet_kg": None,
				"raw_material_weight_kg": None,
			}
		)

		cells = build_cell_map(layout)

		self.assertEqual(cells["K8"], "")
		self.assertEqual(cells["K9"], "")
		self.assertEqual(cells["K10"], "")
		self.assertEqual(cells["K15"], "")
		self.assertEqual(cells["K16"], "")
		self.assertEqual(cells["T6"], "")
		self.assertEqual(cells["T8"], "")
		self.assertEqual(cells["U8"], "")
		for value in cells.values():
			self.assertNotEqual(value, "None")


class TestWorkbookRendering(SheetCuttingLayoutTestCase):
	def test_header_settings_populate_document_cells_and_logo(self) -> None:
		from sheet_cutting_layout.services.export_service import build_multi_sheet_workbook

		with tempfile.NamedTemporaryFile(suffix=".png") as logo_file:
			logo_file.write(
				base64.b64decode(
					"iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGA"
					"WjR9awAAAABJRU5ErkJggg=="
				)
			)
			logo_file.flush()

			workbook = build_multi_sheet_workbook(
				[("Sheet", _base_layout())],
				header_settings={
					"logo_path": logo_file.name,
					"document_number": "SCL/DOC/09",
					"revision_number": "04",
					"revision_date": "2026-06-23",
					"page_text": "01 OF 02",
				},
			)
			stream = io.BytesIO()
			workbook.save(stream)

		worksheet = load_workbook(io.BytesIO(stream.getvalue())).active
		self.assertEqual(worksheet["Q1"].value, "DOC. NO.: SCL/DOC/09")
		self.assertEqual(worksheet["Q2"].value, "REV. NO.: 04")
		self.assertEqual(worksheet["Q3"].value, "REV DATE.: 23.06.2026")
		self.assertEqual(worksheet["Q4"].value, "PAGE: 01 OF 02")
		self.assertEqual(len(getattr(worksheet, "_images", [])), 1)

	def test_renders_cells_into_template_and_returns_xlsx_bytes(self) -> None:
		layout = _base_layout()
		layout.update(
			{
				"is_lh_rh": True,
				"part_names": ["Brkt bumper top", "Brkt bumper top RH"],
				"part_numbers": ["0102AAG06400SHR", "0102AAG06410SHR"],
			}
		)

		content = _render_workbook_content(layout)

		self.assertIsInstance(content, bytes)
		self.assertGreater(len(content), 0)
		self.assertEqual(content[:2], b"PK")

		workbook = load_workbook(io.BytesIO(content))
		worksheet = workbook.active
		self.assertEqual(worksheet["B1"].value, "Acme Press Parts")
		self.assertEqual(worksheet["A5"].value, "Sheet Cutting Layout No:- SCL-EXPORT-LAYOUT-001")
		self.assertEqual(worksheet["G5"].value, "Part Name:-Brkt bumper top & Brkt bumper top RH")
		self.assertEqual(worksheet["N5"].value, "Part Number:-0102AAG06400SHR/0102AAG06410SHR")
		self.assertEqual(worksheet["T6"].value, 31.44)
		self.assertEqual(worksheet["K10"].value, 15.72)
		self.assertEqual(worksheet["K14"].value, 2)
		self.assertEqual(worksheet["Q1"].value, "DOC. NO.:")

	def test_empty_string_cells_clear_the_target_cell(self) -> None:
		layout = _base_layout()
		layout["sheet_thickness_mm"] = None

		content = _render_workbook_content(layout)
		workbook = load_workbook(io.BytesIO(content))
		worksheet = workbook.active
		self.assertIn(worksheet["K8"].value, (None, ""))
		self.assertIn(worksheet["O9"].value, (None, ""))
		self.assertIn(worksheet["Q9"].value, (None, ""))
		self.assertIn(worksheet["U11"].value, (None, ""))
		for coordinate in ("K22", "L22", "M22"):
			self.assertIn(worksheet[coordinate].value, (None, ""))

	def test_formula_like_text_cells_are_escaped(self) -> None:
		layout = _base_layout()
		layout.update(
			{
				"company": "=2+2",
				"part_name": "+part",
				"part_numbers": ["-PART001SHR"],
				"project_name": "@project",
				"end_pieces": [{"disposition": "Reuse", "used_for_finished_part": "=EP-001"}],
			}
		)

		content = _render_workbook_content(layout)
		workbook = load_workbook(io.BytesIO(content))
		worksheet = workbook.active

		self.assertEqual(worksheet["B1"].value, "'=2+2")
		self.assertEqual(worksheet["G5"].value, "Part Name:-+part")
		self.assertEqual(worksheet["N5"].value, "Part Number:--PART001SHR")
		self.assertEqual(worksheet["B6"].value, "'@project")
		self.assertEqual(worksheet["O9"].value, "'=EP-001")
