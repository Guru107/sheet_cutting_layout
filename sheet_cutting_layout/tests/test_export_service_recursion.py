from __future__ import annotations

from collections.abc import Callable

from openpyxl import load_workbook

from sheet_cutting_layout.tests.base import SheetCuttingLayoutTestCase


class _FakeEndPiece:
	def __init__(self, disposition: str = "Scrap", child_layout: str | None = None) -> None:
		self.disposition = disposition
		self.child_layout = child_layout


class _FakeLayout:
	def __init__(self, name: str, end_pieces: list[_FakeEndPiece] | None = None) -> None:
		self.name = name
		self.end_pieces = end_pieces or []


def _base_page(
	name: str,
	*,
	part_name: str = "Bracket",
	part_numbers: list[str] | None = None,
) -> tuple[str, dict[str, object]]:
	return (
		name,
		{
			"company": "Test Co",
			"part_name": part_name,
			"part_numbers": part_numbers or ["P-001"],
			"is_lh_rh": False,
			"project": "PRJ",
			"project_name": "Project",
			"parts_per_sheet": 2,
			"end_pieces": [],
		},
	)


class TestWalkLayoutTree(SheetCuttingLayoutTestCase):
	def _registry(self, *layouts: _FakeLayout) -> Callable[[str], _FakeLayout]:
		registry = {layout.name: layout for layout in layouts}
		return lambda name: registry[name]

	def test_layout_without_children_yields_only_parent(self) -> None:
		parent = _FakeLayout(
			"L1",
			end_pieces=[
				_FakeEndPiece(disposition="Scrap"),
				_FakeEndPiece(disposition="Reuse"),
			],
		)

		from sheet_cutting_layout.services.export_service import walk_layout_tree

		pages = walk_layout_tree(parent, self._registry(parent))

		self.assertEqual([page.name for page in pages], ["L1"])

	def test_single_child_layout_appended_after_parent(self) -> None:
		child = _FakeLayout("L2")
		parent = _FakeLayout("L1", end_pieces=[_FakeEndPiece("Reuse", "L2")])

		from sheet_cutting_layout.services.export_service import walk_layout_tree

		pages = walk_layout_tree(parent, self._registry(parent, child))

		self.assertEqual([page.name for page in pages], ["L1", "L2"])

	def test_two_level_recursion_is_depth_first(self) -> None:
		grandchild = _FakeLayout("L3")
		child = _FakeLayout("L2", end_pieces=[_FakeEndPiece("Reuse", "L3")])
		parent = _FakeLayout("L1", end_pieces=[_FakeEndPiece("Reuse", "L2")])

		from sheet_cutting_layout.services.export_service import walk_layout_tree

		pages = walk_layout_tree(parent, self._registry(parent, child, grandchild))

		self.assertEqual([page.name for page in pages], ["L1", "L2", "L3"])

	def test_two_children_preserve_end_piece_order(self) -> None:
		child_a = _FakeLayout("L2A")
		child_b = _FakeLayout("L2B")
		parent = _FakeLayout(
			"L1",
			end_pieces=[
				_FakeEndPiece("Reuse", "L2A"),
				_FakeEndPiece("Reuse", "L2B"),
			],
		)

		from sheet_cutting_layout.services.export_service import walk_layout_tree

		pages = walk_layout_tree(parent, self._registry(parent, child_a, child_b))

		self.assertEqual([page.name for page in pages], ["L1", "L2A", "L2B"])

	def test_visited_guard_breaks_a_cycle(self) -> None:
		parent = _FakeLayout("L1")
		child = _FakeLayout("L2", end_pieces=[_FakeEndPiece("Reuse", "L1")])
		parent.end_pieces = [_FakeEndPiece("Reuse", "L2")]

		from sheet_cutting_layout.services.export_service import walk_layout_tree

		pages = walk_layout_tree(parent, self._registry(parent, child))

		self.assertEqual([page.name for page in pages], ["L1", "L2"])


class TestUniqueSheetTitle(SheetCuttingLayoutTestCase):
	def test_plain_title_passes_through(self) -> None:
		from sheet_cutting_layout.services.export_service import unique_sheet_title

		self.assertEqual(unique_sheet_title("L1", set()), "L1")

	def test_invalid_characters_are_replaced_with_underscore(self) -> None:
		from sheet_cutting_layout.services.export_service import unique_sheet_title

		title = unique_sheet_title(r"SCL/0102:AAG[06400]\X*?", set())

		for forbidden in "\\/*?:[]":
			self.assertNotIn(forbidden, title)

	def test_title_is_truncated_to_31_characters(self) -> None:
		from sheet_cutting_layout.services.export_service import unique_sheet_title

		title = unique_sheet_title("A" * 50, set())

		self.assertLessEqual(len(title), 31)
		self.assertEqual(title, "A" * 31)

	def test_duplicate_titles_get_a_numeric_suffix(self) -> None:
		from sheet_cutting_layout.services.export_service import unique_sheet_title

		self.assertEqual(unique_sheet_title("L1", {"L1"}), "L1 (2)")

	def test_suffix_keeps_title_within_31_characters(self) -> None:
		from sheet_cutting_layout.services.export_service import unique_sheet_title

		title = unique_sheet_title("A" * 31, {"A" * 31})

		self.assertLessEqual(len(title), 31)
		self.assertNotEqual(title, "A" * 31)
		self.assertTrue(title.endswith("(2)"))

	def test_empty_base_falls_back_to_sheet(self) -> None:
		from sheet_cutting_layout.services.export_service import unique_sheet_title

		self.assertEqual(unique_sheet_title("", set()), "Sheet")


class TestBuildMultiSheetWorkbook(SheetCuttingLayoutTestCase):
	def test_single_page_workbook_has_one_sheet(self) -> None:
		from sheet_cutting_layout.services.export_service import build_multi_sheet_workbook

		workbook = build_multi_sheet_workbook([_base_page("L1")])

		self.assertEqual(workbook.sheetnames, ["L1"])

	def test_one_worksheet_per_page_in_order(self) -> None:
		from sheet_cutting_layout.services.export_service import build_multi_sheet_workbook

		workbook = build_multi_sheet_workbook([_base_page("L1"), _base_page("L2"), _base_page("L3")])

		self.assertEqual(workbook.sheetnames, ["L1", "L2", "L3"])

	def test_duplicate_page_titles_are_made_unique(self) -> None:
		from sheet_cutting_layout.services.export_service import build_multi_sheet_workbook

		workbook = build_multi_sheet_workbook([_base_page("DUP"), _base_page("DUP")])

		self.assertEqual(workbook.sheetnames, ["DUP", "DUP (2)"])

	def test_each_sheet_carries_its_own_part_name(self) -> None:
		from sheet_cutting_layout.services.export_service import build_multi_sheet_workbook

		workbook = build_multi_sheet_workbook(
			[
				_base_page("L1", part_name="Parent Bracket"),
				_base_page("L2", part_name="End Piece Bracket"),
			]
		)

		self.assertEqual(workbook["L1"]["G5"].value, "Part Name:-Parent Bracket")
		self.assertEqual(workbook["L2"]["G5"].value, "Part Name:-End Piece Bracket")

	def test_cloned_sheets_preserve_approved_template_format(self) -> None:
		from sheet_cutting_layout.services.export_service import build_multi_sheet_workbook, default_template_path

		workbook = build_multi_sheet_workbook([_base_page("L1"), _base_page("L2")])
		template_image_count = len(load_workbook(default_template_path()).active._images)

		self.assertGreaterEqual(template_image_count, 2)

		for sheet_name in ("L1", "L2"):
			worksheet = workbook[sheet_name]
			merged_ranges = {str(merged_range) for merged_range in worksheet.merged_cells.ranges}
			self.assertIn("G5:M5", merged_ranges)
			self.assertIn("N5:Q5", merged_ranges)
			self.assertIn("R38:U40", merged_ranges)
			self.assertEqual(worksheet.page_setup.orientation, "landscape")
			self.assertEqual(worksheet["R38"].value, "Released By    \nManagement Rep")
			self.assertEqual(len(worksheet._images), template_image_count)

	def test_empty_page_list_raises(self) -> None:
		from sheet_cutting_layout.services.export_service import build_multi_sheet_workbook

		with self.assertRaises(ValueError):
			build_multi_sheet_workbook([])
