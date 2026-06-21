from __future__ import annotations

from io import BytesIO

import frappe
from openpyxl import load_workbook

from sheet_cutting_layout.sheet_cutting_layout.doctype.sheet_cutting_layout.sheet_cutting_layout import (
	download_sheet_cutting_layout,
)
from sheet_cutting_layout.tests.base import SheetCuttingLayoutTestCase
from sheet_cutting_layout.tests.factories import ensure_item, ensure_project

_PART_NUMBER_PREFIX = "Part Number:-"


def _part_number(value: object) -> str:
	return str(value or "").removeprefix(_PART_NUMBER_PREFIX)


def _release(layout_name: str) -> None:
	layout = frappe.get_doc("Sheet Cutting Layout", layout_name)
	layout.db_set("status", "Approved by Purchase", update_modified=False)
	layout.reload()
	layout.status = "Released"
	layout.submit()


class _RecursiveExportFixtureMixin:
	suffix: str
	project: str
	raw_material: str
	scrap_item: str
	parent_part: str
	child_part: str
	end_piece_item: str
	child: str
	parent: str

	@classmethod
	def _strip_fields(cls) -> dict[str, object]:
		return {
			"sheet_thickness_mm": 2,
			"sheet_width_mm": 1000,
			"sheet_length_mm": 1000,
			"strip_thickness_mm": 2,
			"strip_width_mm": 500,
			"strip_length_mm": 1000,
			"parts_per_strip": 1,
			"no_of_strips": 1,
		}

	@classmethod
	def _setup_items(cls) -> None:
		cls.project = ensure_project()
		cls.raw_material = ensure_item(f"SCLXRM{cls.suffix}", stock_uom="Kg")
		cls.scrap_item = ensure_item(f"SCLXSCRAP{cls.suffix}", stock_uom="Kg")
		cls.parent_part = ensure_item(f"SCLXPARENT{cls.suffix}SHR", stock_uom="Nos")
		cls.child_part = ensure_item(f"SCLXCHILD{cls.suffix}SHR", stock_uom="Nos")
		cls.end_piece_item = ensure_item(f"{cls.parent_part}-EP-2x500x1000", stock_uom="Kg")

	@classmethod
	def _make_child_layout(cls) -> str:
		doc = frappe.get_doc(
			{
				"doctype": "Sheet Cutting Layout",
				"layout_code": f"SCL-X-CHILD-{cls.suffix}",
				"project": cls.project,
				"revision_no": 1,
				"status": "Draft",
				"raw_material_item": cls.end_piece_item,
				"process_scrap_item": cls.scrap_item,
				"finished_part_code": cls.child_part,
				"net_weight_per_part_kg": 7.86,
				**cls._strip_fields(),
				"sheet_width_mm": 500,
				"sheet_length_mm": 1000,
			}
		)
		doc.insert(ignore_permissions=True)
		return doc.name

	@classmethod
	def _parent_doc_values(cls, **overrides: object) -> dict[str, object]:
		values: dict[str, object] = {
			"doctype": "Sheet Cutting Layout",
			"layout_code": f"SCL-X-PARENT-{cls.suffix}",
			"project": cls.project,
			"revision_no": 1,
			"status": "Draft",
			"raw_material_item": cls.raw_material,
			"process_scrap_item": cls.scrap_item,
			"finished_part_code": cls.parent_part,
			"net_weight_per_part_kg": 7.86,
			"end_pieces": [
				{
					"doctype": "Layout End Piece",
					"width_mm": 500,
					"length_mm": 1000,
					"disposition": "Reuse",
					"used_for_finished_part": cls.child_part,
					"scrap_item": cls.scrap_item,
					"child_layout": cls.child,
					"bom_quantity": 1,
					"net_weight_per_part_kg": 7.86,
				}
			],
			**cls._strip_fields(),
		}
		values.update(overrides)
		return values

	@classmethod
	def _make_parent_layout(cls, **overrides: object) -> str:
		doc = frappe.get_doc(cls._parent_doc_values(**overrides))
		doc.insert(ignore_permissions=True)
		return doc.name

	def _download_workbook(self):
		frappe.response.clear()
		download_sheet_cutting_layout(self.parent)
		return load_workbook(BytesIO(frappe.response["filecontent"]))


class TestDownloadRecursiveLayout(_RecursiveExportFixtureMixin, SheetCuttingLayoutTestCase):
	@classmethod
	def setUpClass(cls) -> None:
		super().setUpClass()
		cls.suffix = frappe.generate_hash(length=8).upper()
		cls._setup_items()
		cls.child = cls._make_child_layout()
		cls.parent = cls._make_parent_layout()

	def test_download_returns_xlsx_response(self) -> None:
		frappe.response.clear()

		download_sheet_cutting_layout(self.parent)

		self.assertEqual(frappe.response["type"], "binary")
		self.assertTrue(str(frappe.response["filename"]).endswith(".xlsx"))
		self.assertEqual(frappe.response["filecontent"][:2], b"PK")

	def test_workbook_has_one_sheet_per_layout_in_tree(self) -> None:
		workbook = self._download_workbook()

		self.assertEqual(len(workbook.sheetnames), 2)

	def test_parent_and_child_sheets_carry_distinct_part_numbers(self) -> None:
		workbook = self._download_workbook()

		part_numbers = {_part_number(workbook[name]["N5"].value) for name in workbook.sheetnames}
		self.assertIn(self.parent_part, part_numbers)
		self.assertIn(self.child_part, part_numbers)


class TestDownloadReleasedRecursiveLhRhLayout(_RecursiveExportFixtureMixin, SheetCuttingLayoutTestCase):
	@classmethod
	def setUpClass(cls) -> None:
		super().setUpClass()
		cls.suffix = frappe.generate_hash(length=8).upper()
		cls._setup_items()
		cls.twin_part = ensure_item(f"SCLXTWIN{cls.suffix}SHR", stock_uom="Nos")
		cls.end_piece_item = ensure_item(f"{cls.parent_part}-{cls.twin_part}-EP-2x500x1000", stock_uom="Kg")
		cls.child = cls._make_child_layout()
		cls.parent = cls._make_parent_layout(
			layout_code=f"SCL-X-LHRH-{cls.suffix}",
			is_lh_rh=1,
			orientation="LH",
			twin_finished_part=cls.twin_part,
		)
		_release(cls.child)
		_release(cls.parent)

	def test_released_parent_sheet_shows_joined_lh_rh_part_number(self) -> None:
		workbook = self._download_workbook()

		joined = _part_number(workbook[workbook.sheetnames[0]]["N5"].value)

		self.assertIn(self.parent_part, joined)
		self.assertIn(self.twin_part, joined)
		self.assertIn("/", joined)
		self.assertNotIn("_", joined)

	def test_released_workbook_has_parent_and_child_sheets(self) -> None:
		workbook = self._download_workbook()

		part_numbers = {_part_number(workbook[name]["N5"].value) for name in workbook.sheetnames}

		self.assertEqual(len(workbook.sheetnames), 2)
		self.assertTrue(any(self.child_part in value for value in part_numbers))
